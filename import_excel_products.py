#!/usr/bin/env python3
"""Import products from 2026-10-09商品.xlsx into catalog.js + item pages + images."""
from __future__ import annotations

import json
import re
import time
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from html import escape
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "2026-10-09商品.xlsx"
CATALOG = ROOT / "js" / "catalog.js"
IMG_DIR = (ROOT / "img" / "products").resolve()
ORIGIN = "https://hipobuyqcsheets.com"
INVITE = "R71GHKM1I"
CNY_TO_USD = 7.2
WORKERS = 10
UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"

CAT_LABELS = {
    "shoes": "Shoes",
    "t-shirts": "T-Shirts",
    "hoodies": "Hoodies",
    "jackets": "Jackets",
    "pants": "Pants / Shorts",
    "sets": "Sets",
    "jersey": "Jerseys",
    "underwear": "Underwear",
    "headwear": "Headwear",
    "watches": "Watches",
    "accessories": "Accessories",
    "other": "Other",
    "glasses": "Glasses",
    "perfume": "Perfume",
    "bricks": "Building Blocks",
}

BRAND_HINTS = {
    "nike", "adidas", "jordan", "new balance", "asics", "salomon", "gucci", "lv",
    "louis vuitton", "dior", "chanel", "prada", "balenciaga", "chrome hearts",
    "ralph lauren", "stone island", "cartier", "apple", "hermes", "hermès",
    "bape", "supreme", "stussy", "essentials", "moncler", "canada goose",
    "the north face", "carhartt", "off-white", "amiri", "corteiz", "trapstar",
    "denim tears", "yeezy", "ugg", "mlb", "new era", "ray-ban", "casio",
    "rolex", "swatch", "sony", "fendi", "burberry", "coach", "goyard",
    "arc'teryx", "arcteryx", "patagonia", "hellstar", "nocta", "skims",
}


def slugify(text: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", (text or "").lower())
    s = re.sub(r"^-+|-+$", "", s)
    s = re.sub(r"-+", "-", s)
    return s[:60]


def item_slug(p: dict) -> str:
    num = re.sub(r"^p-", "", str(p.get("id") or ""), count=1)
    base = slugify(str(p.get("title") or ""))
    return f"{base}-{num}" if base else num


def load_catalog() -> list[dict]:
    raw = CATALOG.read_text(encoding="utf-8")
    start = raw.index("KF.products = ") + len("KF.products = ")
    products, _ = json.JSONDecoder().raw_decode(raw[start:])
    return products


def write_catalog(products: list[dict]) -> None:
    body = json.dumps(products, ensure_ascii=False, indent=2)
    CATALOG.write_text(f"KF.products = {body};\n", encoding="utf-8")


def map_category(tags: list[str], title: str) -> str:
    blob = " ".join(tags + [title]).lower()
    rules = [
        (r"shoe|sneaker|slide|boot|dunk|jordan|yeezy|asics|trainer", "shoes"),
        (r"hoodie|sweatshirt|sweater", "hoodies"),
        (r"jacket|puffer|parka|coat|vest|nuptse", "jackets"),
        (r"t-?shirt|tee\b", "t-shirts"),
        (r"jersey", "jersey"),
        (r"short|pants|jean|cargo|bottom", "pants"),
        (r"cap|hat|beanie|headwear", "headwear"),
        (r"watch|rolex|g-?shock", "watches"),
        (r"suit|tracksuit|set\b", "sets"),
        (r"sunglass|glass|belt|bag|wallet|accessori|perfume|jewelry|pendant", "accessories"),
        (r"airpod|electronic|headphone|earbud", "other"),
    ]
    for pat, cat in rules:
        if re.search(pat, blob):
            return cat
    return "other"


def map_collection(tags: list[str], title: str) -> str:
    for t in tags:
        low = t.strip().lower()
        if low in {"top trending", "shoes", "bag", "shorts", "pants", "jeans", "hoodies",
                   "jacket", "vest", "suit", "dress", "accessories", "electronics",
                   "sunglasses", "slides", "belt", "cap", "t-shirt"}:
            continue
        if low in BRAND_HINTS or any(b in low for b in BRAND_HINTS):
            return t.strip()
    title_l = title.lower()
    for b in sorted(BRAND_HINTS, key=len, reverse=True):
        if b in title_l:
            return b.title() if b.islower() else b
    return tags[0].strip() if tags else "HipoBuySpreadsheet"


def sizes_for(cat: str) -> list[str]:
    if cat == "shoes":
        return ["40", "41", "42", "43", "44"]
    if cat in {"watches", "accessories", "other", "glasses", "perfume"}:
        return ["One Size"]
    return ["S", "M", "L", "XL"]


def fetch_weidian(item_id: str) -> dict:
    url = f'https://thor.weidian.com/detail/getItemSkuInfo/1.0?param={{"itemId":"{item_id}"}}'
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Referer": "https://weidian.com/"})
    with urllib.request.urlopen(req, timeout=20) as resp:
        data = json.load(resp)
    if not isinstance(data, dict) or (data.get("status") or {}).get("code") != 0:
        raise RuntimeError(f"weidian status {data.get('status')}")
    return data.get("result") or {}


def cny_fen_to_usd(fen) -> float:
    try:
        cny = float(fen) / 100.0
    except (TypeError, ValueError):
        return 0.0
    return round(cny / CNY_TO_USD, 2)


def download_webp(img_url: str, dest: Path) -> bool:
    if dest.exists() and dest.stat().st_size > 1000:
        return True
    if not img_url:
        return False
    req = urllib.request.Request(img_url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=30) as resp:
        raw = resp.read()
    img = Image.open(BytesIO(raw)).convert("RGB")
    dest.parent.mkdir(parents=True, exist_ok=True)
    img.save(dest, "WEBP", quality=82, method=4)
    return True


def enrich(row: dict) -> dict:
    item_id = row["id"]
    out = dict(row)
    out["ok"] = False
    out["price"] = 0.0
    out["image_url"] = ""
    out["api_title"] = ""
    for attempt in range(3):
        try:
            res = fetch_weidian(item_id)
            fen = res.get("itemDiscountLowPrice") or res.get("itemOriginalLowPrice") or 0
            out["price"] = cny_fen_to_usd(fen)
            out["image_url"] = str(res.get("itemMainPic") or "")
            out["api_title"] = str(res.get("itemTitle") or "")
            if out["image_url"]:
                download_webp(out["image_url"], IMG_DIR / f"{item_id}.webp")
            out["ok"] = True
            return out
        except Exception as e:
            out["error"] = str(e)
            time.sleep(0.4 * (attempt + 1))
    return out


def read_excel() -> list[dict]:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb.active
    rows = []
    for name, url, tags in list(ws.iter_rows(values_only=True))[1:]:
        if not url:
            continue
        m = re.search(r"itemID=(\d+)", str(url), re.I)
        if not m:
            continue
        tag_list = [t.strip() for t in str(tags or "").split(",") if t.strip()]
        title = str(name or "").strip() or f"Item {m.group(1)}"
        rows.append({
            "id": m.group(1),
            "title": title,
            "url": str(url).strip(),
            "tags": tag_list,
            "trending": any(t.lower() == "top trending" for t in tag_list),
        })
    return rows


def product_from_row(row: dict, featured: bool) -> dict:
    cat = map_category(row["tags"], row["title"])
    collection = map_collection(row["tags"], row["title"])
    img = f"img/products/{row['id']}.webp"
    return {
        "id": f"p-{row['id']}",
        "title": row["title"],
        "price": float(row.get("price") or 0),
        "rating": 4.7,
        "category": cat,
        "collection": collection,
        "qc": True,
        "featured": featured,
        "seller": "wen",
        "source": "Weidian",
        "sourceUrl": f"https://weidian.com/item.html?itemID={row['id']}",
        "image": img,
        "gallery": [img],
        "sizes": sizes_for(cat),
        "summary": f"{collection} · {', '.join(row['tags'][:4])}" if row["tags"] else collection,
    }


def hipobuy_url(source_url: str) -> str:
    m = re.search(r"itemID=(\d+)", source_url or "", re.I)
    if not m:
        return source_url or "https://hipobuy.com/"
    pid = m.group(1)
    raw = f"https://weidian.com/item.html?itemID={pid}"
    from urllib.parse import quote
    enc = quote(raw, safe="")
    return (
        f"https://hipobuy.com/product/weidian/{pid}"
        f"?source=SEARCH_LIST&keyword={enc}&inviteCode={INVITE}"
    )


def product_card_html(p: dict, depth: int = 2) -> str:
    slug = item_slug(p)
    prefix = "../" * depth
    rel = f"../{slug}/" if depth == 2 else f"item/{slug}/"
    # related cards on item pages use sibling relative links
    href = f"../{slug}/"
    img = str(p.get("image") or "")
    img_src = prefix + img if not img.startswith("http") else img
    # for related inside item/xxx/: ../../img/... and ../other-slug/
    if depth == 2:
        img_src = "../../" + img.lstrip("./")
        href = f"../{slug}/"
    title = escape(str(p.get("title") or ""))
    price = float(p.get("price") or 0)
    return (
        f'<a class="product-card" href="{escape(href)}">'
        f'<div class="thumb"><img src="{escape(img_src)}" alt="{title}" loading="lazy" decoding="async" /></div>'
        f"<h3>{title}</h3><b>${price:.2f}</b></a>"
    )


def item_page_html(p: dict, related: list[dict]) -> str:
    slug = item_slug(p)
    path = f"/item/{slug}/"
    canonical = ORIGIN + path
    name = str(p.get("title") or "Find")
    page_title = f"{name} — HipoBuySpreadsheet"
    price = float(p.get("price") or 0)
    cat = str(p.get("category") or "")
    label = CAT_LABELS.get(cat, cat.replace("-", " ").title() or "Find")
    source = str(p.get("source") or "Weidian")
    img_rel = "../../" + str(p.get("image") or "").lstrip("./")
    img_abs = ORIGIN + "/" + str(p.get("image") or "").lstrip("./")
    desc = (
        f"{name} on the HipoBuySpreadsheet. {label} find from {source}, "
        f"listed at ${price:.2f}. Preview QC photos here, then open Hipobuy to order."
    )
    buy = hipobuy_url(str(p.get("sourceUrl") or ""))
    brand = str(p.get("collection") or "HipoBuySpreadsheet")
    ld = {
        "@context": "https://schema.org",
        "@type": "Product",
        "name": name,
        "image": img_abs,
        "description": desc,
        "brand": {"@type": "Brand", "name": brand},
        "offers": {
            "@type": "Offer",
            "price": f"{price:.2f}",
            "priceCurrency": "USD",
            "availability": "https://schema.org/InStock",
            "url": canonical,
        },
    }
    rel = "\n        ".join(product_card_html(r, depth=2) for r in related[:8])
    if not rel:
        rel = "<p>More finds in the spreadsheet.</p>"
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>{escape(page_title)}</title>
  <meta name="description" content="{escape(desc)}" />
  <meta name="robots" content="index, follow, max-image-preview:large" />
  <link rel="canonical" href="{escape(canonical)}" />
  <meta property="og:type" content="product" />
  <meta property="og:site_name" content="hipobuyqcsheets" />
  <meta property="og:title" content="{escape(page_title)}" />
  <meta property="og:description" content="{escape(desc)}" />
  <meta property="og:url" content="{escape(canonical)}" />
  <meta property="og:image" content="{escape(img_abs)}" />
  <meta property="product:price:amount" content="{price:.2f}" />
  <meta property="product:price:currency" content="USD" />
  <script type="application/ld+json">{json.dumps(ld, ensure_ascii=False)}</script>
  <link rel="icon" href="../../img/logo.png" type="image/png" />
  <link rel="preconnect" href="https://fonts.googleapis.com" />
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin />
  <link href="https://fonts.googleapis.com/css2?family=DM+Sans:opsz,wght@9..40,400;9..40,500;9..40,600;9..40,700&display=swap" rel="stylesheet" />
  <link rel="stylesheet" href="../../css/home.css" />
  <link rel="stylesheet" href="../../css/pages.css" />
  <link rel="stylesheet" href="../../css/styles.css" />
</head>
<body data-page="item" data-item-id="{escape(str(p.get('id') or ''))}">
  <header class="site-header">
    <div class="topbar">
      <div class="wrap topbar-row">
        <a class="brand" href="../../index.html" aria-label="HipoBuySpreadsheet home">
          <img src="../../img/logo.png" alt="" width="40" height="40" />
          <span class="brand-text">
            <span class="brand-name">HipoBuy<span>Spreadsheet</span></span>
            <span class="brand-sub">QC Sheet · Finder · Catalog</span>
          </span>
        </a>
        <nav class="nav-rail" aria-label="Primary">
          <a href="../../index.html">Home</a>
          <a href="../../finder.html">QC Finder</a>
          <a href="../../sheets.html">QC Sheets</a>
          <a href="../../browse.html">Spreadsheet</a>
          <a href="../../guides/index.html">Guides</a>
        </nav>
        <div class="top-actions">
          <a class="top-cta" href="../../finder.html">Open Finder</a>
          <button class="menu-btn" type="button" aria-label="Menu" aria-expanded="false" onclick="document.body.classList.toggle('nav-open');this.setAttribute('aria-expanded',document.body.classList.contains('nav-open'))">
            <span></span><span></span>
          </button>
        </div>
      </div>
    </div>
    <div class="cat-strip">
      <div class="wrap cat-row">
        <a href="../../browse.html?cat=shoes">Sneakers</a>
        <a href="../../browse.html?cat=hoodies">Hoodies</a>
        <a href="../../browse.html?cat=t-shirts">Tees</a>
        <a href="../../browse.html?cat=jackets">Jackets</a>
        <a href="../../browse.html?cat=jersey">Jerseys</a>
        <a href="../../browse.html?cat=accessories">Accessories</a>
        <a href="../../browse.html?cat=watches">Watches</a>
        <a href="../../browse.html?cat=headwear">Hats</a>
        <a href="../../browse.html">All finds →</a>
      </div>
    </div>
  </header>
  <main class="wrap item-page">
    <nav class="crumbs">
      <a href="../../index.html">Home</a><span>/</span>
      <a href="../../browse.html">Shop</a><span>/</span>
      <a href="../../browse.html?cat={escape(cat)}">{escape(label)}</a><span>/</span>
      <span>{escape(name)}</span>
    </nav>
    <div class="item-layout">
      <div class="gallery">
        <div class="main">
          <img src="{escape(img_rel)}" alt="{escape(name)}" />
        </div>
      </div>
      <div class="item-info">
        <h1>{escape(name)}</h1>
        <div class="price">${price:.2f}</div>
        <p class="item-rating">QC photos on this find · {escape(source)}</p>
        <a class="btn btn-gold btn-buy" href="{escape(buy)}" target="_blank" rel="noopener">Buy on Hipobuy</a>
        <p class="item-note">{escape(desc)}</p>
      </div>
    </div>
    <section class="related">
      <h2>More {escape(label)} finds</h2>
      <div class="product-grid">
        {rel}
      </div>
    </section>
  </main>
  <footer class="site-footer">
    <div class="wrap foot-note">© 2026 HipoBuySpreadsheet · hipobuyqcsheets.com — independent directory.</div>
  </footer>
</body>
</html>
"""


def append_sitemap(slugs: list[str]) -> None:
    path = ROOT / "sitemap.xml"
    if not path.exists():
        return
    text = path.read_text(encoding="utf-8")
    today = "2026-10-09"
    entries = []
    for s in slugs:
        loc = f"{ORIGIN}/item/{s}/"
        if loc in text:
            continue
        entries.append(
            f"  <url><loc>{loc}</loc><lastmod>{today}</lastmod>"
            f"<changefreq>weekly</changefreq><priority>0.6</priority></url>"
        )
    if not entries:
        return
    text = text.replace("</urlset>", "\n".join(entries) + "\n</urlset>")
    path.write_text(text, encoding="utf-8")


def main() -> None:
    print("Reading Excel…")
    rows = read_excel()
    existing = load_catalog()
    existing_ids = {re.sub(r"^p-", "", str(p.get("id") or "")) for p in existing}
    todo = [r for r in rows if r["id"] not in existing_ids]
    print(f"Excel {len(rows)} · catalog {len(existing)} · new {len(todo)}")
    if not todo:
        print("Nothing to add.")
        return

    IMG_DIR.mkdir(parents=True, exist_ok=True)
    enriched = []
    done = 0
    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        futs = {pool.submit(enrich, r): r["id"] for r in todo}
        for fut in as_completed(futs):
            enriched.append(fut.result())
            done += 1
            if done % 25 == 0 or done == len(todo):
                ok = sum(1 for x in enriched if x.get("ok"))
                print(f"  fetched {done}/{len(todo)} (ok {ok})")

    # preserve Excel order
    by_id = {x["id"]: x for x in enriched}
    ordered = [by_id[r["id"]] for r in todo if r["id"] in by_id]

    new_products = []
    for i, row in enumerate(ordered):
        featured = bool(row.get("trending")) or i < 12
        new_products.append(product_from_row(row, featured=featured))

    # newest first in catalog
    merged = new_products + existing
    print(f"Writing catalog.js ({len(merged)} products)…")
    write_catalog(merged)

    # related pool: new + a slice of existing by category
    by_cat: dict[str, list] = {}
    for p in merged:
        by_cat.setdefault(str(p.get("category") or "other"), []).append(p)

    slugs = []
    print("Writing item pages…")
    for p in new_products:
        slug = item_slug(p)
        cat = str(p.get("category") or "other")
        related = [x for x in by_cat.get(cat, []) if x.get("id") != p.get("id")][:8]
        folder = ROOT / "item" / slug
        folder.mkdir(parents=True, exist_ok=True)
        (folder / "index.html").write_text(item_page_html(p, related), encoding="utf-8")
        slugs.append(slug)

    append_sitemap(slugs)
    ok_imgs = sum(1 for r in ordered if (IMG_DIR / f"{r['id']}.webp").exists())
    fail = [r for r in ordered if not r.get("ok")]
    print(f"Done. Added {len(new_products)} products, images {ok_imgs}/{len(ordered)}, api fails {len(fail)}")
    if fail[:5]:
        print("Sample fails:", [(f["id"], f.get("error")) for f in fail[:5]])


if __name__ == "__main__":
    main()
